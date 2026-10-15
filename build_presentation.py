import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# --- COLOR PALETTE (Academic Minimal: Warm Ivory, Charcoal, Dusty Burgundy) ---
COLOR_BG = RGBColor(250, 248, 245)          # #FAF8F5 Warm Ivory
COLOR_CARD = RGBColor(255, 255, 255)        # #FFFFFF Pure Clean White
COLOR_CARD_BORDER = RGBColor(226, 219, 209) # #E2DBD1 Subtle Warm Border
COLOR_TEXT_PRIMARY = RGBColor(42, 36, 33)   # #2A2421 Dark Charcoal / Espresso
COLOR_TEXT_BODY = RGBColor(74, 67, 62)      # #4A433E Soft Charcoal
COLOR_TEXT_MUTED = RGBColor(120, 112, 105)  # #787069 Warm Gray
COLOR_ACCENT = RGBColor(124, 45, 55)        # #7C2D37 Dusty Burgundy (Single Muted Accent)
COLOR_ACCENT_TINT = RGBColor(246, 238, 239) # #F6EEEF Very Light Burgundy Tint
COLOR_RULE = RGBColor(216, 208, 198)        # #D8D0C6 Hairline Divider Rule

FONT_TITLE = "Georgia"
FONT_BODY = "Calibri"

def add_card(slide, left, top, width, height, fill_color=COLOR_CARD, border_color=COLOR_CARD_BORDER, border_width=Pt(0.75)):
    """Helper to add a cleanly styled card rectangle."""
    card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = fill_color
    if border_color is not None:
        card.line.color.rgb = border_color
        card.line.width = border_width
    else:
        card.line.color.rgb = fill_color
        card.line.width = Pt(0)
    return card

def set_slide_background(slide):
    """Fills slide background with warm ivory."""
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLOR_BG
    bg.line.color.rgb = COLOR_BG
    bg.line.width = Pt(0)
    return bg

def add_header(slide, eyebrow, title, subtitle):
    """Creates a consistent academic header."""
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.48), Inches(11.733), Inches(1.15))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    # Eyebrow
    p0 = tf.paragraphs[0]
    p0.text = eyebrow.upper()
    p0.font.name = FONT_BODY
    p0.font.size = Pt(10)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_ACCENT
    p0.space_after = Pt(3)

    # Title
    p1 = tf.add_paragraph()
    p1.text = title
    p1.font.name = FONT_TITLE
    p1.font.size = Pt(26)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_TEXT_PRIMARY
    p1.space_after = Pt(2)

    # Subtitle
    p2 = tf.add_paragraph()
    p2.text = subtitle
    p2.font.name = FONT_BODY
    p2.font.size = Pt(13)
    p2.font.italic = True
    p2.font.color.rgb = COLOR_TEXT_MUTED

    # Hairline divider
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.68), Inches(11.733), Inches(0.015))
    rule.fill.solid()
    rule.fill.fore_color.rgb = COLOR_RULE
    rule.line.color.rgb = COLOR_RULE
    rule.line.width = Pt(0)

def add_footer(slide, slide_num):
    """Adds a clean, academic footer."""
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.85), Inches(11.733), Inches(0.01))
    rule.fill.solid()
    rule.fill.fore_color.rgb = COLOR_CARD_BORDER
    rule.line.color.rgb = COLOR_CARD_BORDER
    rule.line.width = Pt(0)

    # Left text
    tb_l = slide.shapes.add_textbox(Inches(0.8), Inches(6.92), Inches(8.0), Inches(0.3))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True
    tf_l.margin_left = tf_l.margin_top = tf_l.margin_right = tf_l.margin_bottom = 0
    p = tf_l.paragraphs[0]
    p.text = "Telephone Conversation — Wole Soyinka | College English Literature Seminar"
    p.font.name = FONT_BODY
    p.font.size = Pt(9.5)
    p.font.color.rgb = COLOR_TEXT_MUTED

    # Right text
    tb_r = slide.shapes.add_textbox(Inches(10.533), Inches(6.92), Inches(2.0), Inches(0.3))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True
    tf_r.margin_left = tf_r.margin_top = tf_r.margin_right = tf_r.margin_bottom = 0
    p_r = tf_r.paragraphs[0]
    p_r.text = f"Slide {slide_num} of 6"
    p_r.alignment = PP_ALIGN.RIGHT
    p_r.font.name = FONT_BODY
    p_r.font.size = Pt(9.5)
    p_r.font.color.rgb = COLOR_TEXT_MUTED

def set_notes(slide, text):
    """Sets speaker notes on slide."""
    slide.notes_slide.notes_text_frame.text = text

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # =========================================================================
    # SLIDE 1: TITLE SLIDE
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1)

    # Left title content area
    tb_title = slide1.shapes.add_textbox(Inches(0.9), Inches(1.15), Inches(6.8), Inches(3.6))
    tf1 = tb_title.text_frame
    tf1.word_wrap = True
    tf1.margin_left = tf1.margin_top = tf1.margin_right = tf1.margin_bottom = 0

    p_eye = tf1.paragraphs[0]
    p_eye.text = "COLLEGE ENGLISH LITERATURE SEMINAR"
    p_eye.font.name = FONT_BODY
    p_eye.font.size = Pt(11)
    p_eye.font.bold = True
    p_eye.font.color.rgb = COLOR_ACCENT
    p_eye.space_after = Pt(14)

    p_t = tf1.add_paragraph()
    p_t.text = "Telephone Conversation"
    p_t.font.name = FONT_TITLE
    p_t.font.size = Pt(38)
    p_t.font.bold = True
    p_t.font.color.rgb = COLOR_TEXT_PRIMARY
    p_t.space_after = Pt(6)

    p_sub = tf1.add_paragraph()
    p_sub.text = "By Wole Soyinka"
    p_sub.font.name = FONT_TITLE
    p_sub.font.size = Pt(21)
    p_sub.font.italic = True
    p_sub.font.color.rgb = COLOR_ACCENT
    p_sub.space_after = Pt(16)

    # Accent rule
    line_div = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(3.28), Inches(2.2), Inches(0.025))
    line_div.fill.solid()
    line_div.fill.fore_color.rgb = COLOR_ACCENT
    line_div.line.color.rgb = COLOR_ACCENT
    line_div.line.width = Pt(0)

    # Descriptive context
    tb_desc = slide1.shapes.add_textbox(Inches(0.9), Inches(3.52), Inches(6.6), Inches(1.05))
    tf_desc = tb_desc.text_frame
    tf_desc.word_wrap = True
    tf_desc.margin_left = tf_desc.margin_top = tf_desc.margin_right = tf_desc.margin_bottom = 0
    p_desc = tf_desc.paragraphs[0]
    p_desc.text = "A close reading of racial prejudice, everyday discrimination, and the defense of human dignity through irony."
    p_desc.font.name = FONT_BODY
    p_desc.font.size = Pt(13.5)
    p_desc.font.color.rgb = COLOR_TEXT_BODY

    # Student metadata card
    add_card(slide1, Inches(0.9), Inches(4.75), Inches(6.6), Inches(1.8), fill_color=COLOR_CARD, border_color=COLOR_CARD_BORDER, border_width=Pt(0.75))

    tb_meta = slide1.shapes.add_textbox(Inches(1.15), Inches(4.9), Inches(6.1), Inches(1.5))
    tf_meta = tb_meta.text_frame
    tf_meta.word_wrap = True
    tf_meta.margin_left = tf_meta.margin_top = tf_meta.margin_right = tf_meta.margin_bottom = 0

    pm1 = tf_meta.paragraphs[0]
    pm1.text = "PRESENTED BY"
    pm1.font.name = FONT_BODY
    pm1.font.size = Pt(9.5)
    pm1.font.bold = True
    pm1.font.color.rgb = COLOR_TEXT_MUTED
    pm1.space_after = Pt(2)

    pm2 = tf_meta.add_paragraph()
    pm2.text = "[Student Name]"
    pm2.font.name = FONT_TITLE
    pm2.font.size = Pt(15)
    pm2.font.bold = True
    pm2.font.color.rgb = COLOR_TEXT_PRIMARY
    pm2.space_after = Pt(4)

    pm3 = tf_meta.add_paragraph()
    pm3.text = "Department of English Literature • College Seminar Series"
    pm3.font.name = FONT_BODY
    pm3.font.size = Pt(11.5)
    pm3.font.color.rgb = COLOR_TEXT_BODY
    pm3.space_after = Pt(2)

    pm4 = tf_meta.add_paragraph()
    pm4.text = "Topic: Modern African Poetry & Post-Colonial Literature"
    pm4.font.name = FONT_BODY
    pm4.font.size = Pt(11)
    pm4.font.italic = True
    pm4.font.color.rgb = COLOR_TEXT_MUTED

    # Right column: Photograph frame
    add_card(slide1, Inches(8.0), Inches(1.15), Inches(4.4), Inches(4.9), fill_color=COLOR_CARD, border_color=COLOR_CARD_BORDER, border_width=Pt(1.0))

    if os.path.exists("assets/phone_cropped.jpg"):
        slide1.shapes.add_picture("assets/phone_cropped.jpg", Inches(8.2), Inches(1.35), Inches(4.0), Inches(3.95))

    tb_cap1 = slide1.shapes.add_textbox(Inches(8.1), Inches(5.45), Inches(4.2), Inches(0.45))
    tf_cap1 = tb_cap1.text_frame
    tf_cap1.word_wrap = True
    tf_cap1.margin_left = tf_cap1.margin_top = tf_cap1.margin_right = tf_cap1.margin_bottom = 0
    p_cap1 = tf_cap1.paragraphs[0]
    p_cap1.text = "A mid-century telephone: bridging communication across racial barriers."
    p_cap1.alignment = PP_ALIGN.CENTER
    p_cap1.font.name = FONT_BODY
    p_cap1.font.size = Pt(10.5)
    p_cap1.font.italic = True
    p_cap1.font.color.rgb = COLOR_TEXT_MUTED

    add_footer(slide1, 1)
    set_notes(slide1,
        "Good morning everyone, and welcome to today's English literature seminar. My name is [Student Name], "
        "and today I will be presenting Wole Soyinka’s celebrated 1962 poem, 'Telephone Conversation.'\n\n"
        "Set in London during the post-war era, this poem captures an encounter that feels both painfully realistic and "
        "masterfully comic: an educated Black African student phoning a British landlady about a rented room. What begins as a "
        "routine housing inquiry quickly spirals into an intense examination of racial prejudice, polite hypocrisy, and the defense of human dignity.\n\n"
        "Over the next few minutes, we will look at Soyinka's background, trace the narrative sequence of the telephone exchange, "
        "examine its major themes, analyze its literary devices, and reflect on its lasting message for our world today. "
        "Let us begin with the author."
    )

    # =========================================================================
    # SLIDE 2: ABOUT THE POET
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2)
    add_header(slide2, "Biographical Context", "About Wole Soyinka", "Nigerian playwright, poet, and champion of human dignity")

    # Left column: Poet portrait frame
    add_card(slide2, Inches(0.8), Inches(1.9), Inches(3.6), Inches(4.75), fill_color=COLOR_CARD, border_color=COLOR_CARD_BORDER, border_width=Pt(0.75))

    if os.path.exists("assets/soyinka_cropped.jpg"):
        slide2.shapes.add_picture("assets/soyinka_cropped.jpg", Inches(0.95), Inches(2.05), Inches(3.3), Inches(3.75))

    tb_pcap = slide2.shapes.add_textbox(Inches(0.95), Inches(5.95), Inches(3.3), Inches(0.6))
    tf_pcap = tb_pcap.text_frame
    tf_pcap.word_wrap = True
    tf_pcap.margin_left = tf_pcap.margin_top = tf_pcap.margin_right = tf_pcap.margin_bottom = 0
    ppc1 = tf_pcap.paragraphs[0]
    ppc1.text = "Wole Soyinka (b. 13 July 1934)"
    ppc1.font.name = FONT_TITLE
    ppc1.font.size = Pt(11)
    ppc1.font.bold = True
    ppc1.font.color.rgb = COLOR_TEXT_PRIMARY
    ppc1.alignment = PP_ALIGN.CENTER
    ppc2 = tf_pcap.add_paragraph()
    ppc2.text = "Nigerian Poet, Playwright & Nobel Laureate"
    ppc2.font.name = FONT_BODY
    ppc2.font.size = Pt(9.5)
    ppc2.font.color.rgb = COLOR_TEXT_MUTED
    ppc2.alignment = PP_ALIGN.CENTER

    # Right column: Biographical points
    right_x = Inches(4.65)
    right_w = Inches(7.88)

    bio_items = [
        (
            "Nigerian Literary Pioneer",
            "Born on 13 July 1934 in Abeokuta, Nigeria. Soyinka studied at University College, Ibadan, and the University of Leeds. He became one of the foremost dramatists, poets, and cultural critics of modern African literature."
        ),
        (
            "Nobel Prize in Literature (1986)",
            "In 1986, Soyinka made history as the first sub-Saharan African writer to receive the Nobel Prize in Literature, commended for fashioning the drama of existence with wide cultural perspective and poetic resonance."
        ),
        (
            "Lifelong Defense of Human Rights",
            "His literary work fearlessly challenges political oppression, corruption, and social injustice. He endured political imprisonment in Nigeria for opposing war and has continually defended individual freedom and dignity."
        ),
        (
            "Context of 'Telephone Conversation'",
            "Written during his years in England in the early 1960s, the poem reflects the pervasive housing discrimination and subtle 'colour bar' faced by African scholars and immigrants in post-war Britain."
        )
    ]

    card_y = Inches(1.9)
    card_h = Inches(1.1)
    for title_txt, body_txt in bio_items:
        add_card(slide2, right_x, card_y, right_w, card_h, fill_color=COLOR_CARD, border_color=COLOR_CARD_BORDER, border_width=Pt(0.75))

        # Left subtle accent bar
        bar = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, right_x, card_y, Inches(0.08), card_h)
        bar.fill.solid()
        bar.fill.fore_color.rgb = COLOR_ACCENT
        bar.line.color.rgb = COLOR_ACCENT
        bar.line.width = Pt(0)

        tb = slide2.shapes.add_textbox(right_x + Inches(0.22), card_y + Inches(0.1), right_w - Inches(0.35), card_h - Inches(0.2))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_bt = tf.paragraphs[0]
        p_bt.text = title_txt
        p_bt.font.name = FONT_TITLE
        p_bt.font.size = Pt(13)
        p_bt.font.bold = True
        p_bt.font.color.rgb = COLOR_ACCENT
        p_bt.space_after = Pt(2)

        p_bb = tf.add_paragraph()
        p_bb.text = body_txt
        p_bb.font.name = FONT_BODY
        p_bb.font.size = Pt(11)
        p_bb.font.color.rgb = COLOR_TEXT_BODY

        card_y += card_h + Inches(0.12)

    add_footer(slide2, 2)
    set_notes(slide2,
        "To appreciate the poem, we first need to understand the poet behind it. Wole Soyinka was born in Nigeria in 1934 and studied both "
        "in Ibadan and at the University of Leeds in the United Kingdom. He is globally renowned as a playwright, novelist, poet, and essayist.\n\n"
        "In 1986, Soyinka made history by becoming the very first sub-Saharan African writer to win the Nobel Prize in Literature. "
        "The Swedish Academy praised him for his wide cultural perspective and poetic depth.\n\n"
        "Beyond his literary genius, Soyinka is known for his moral courage. He has consistently spoken out against tyranny, censorship, "
        "and human rights abuses—even enduring nearly two years of solitary confinement during the Nigerian Civil War for his anti-war advocacy.\n\n"
        "He wrote 'Telephone Conversation' in the early 1960s, drawing directly upon his lived experiences as a young African scholar in London. "
        "At the time, finding housing was a humiliating ordeal for Black students due to widespread racial prejudice among landlords. "
        "Now, let’s examine how this narrative situation unfolds in the poem."
    )

    # =========================================================================
    # SLIDE 3: SUMMARY OF THE POEM
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3)
    add_header(slide3, "Textual Overview", "Summary: The Flow of the Conversation", "A natural sequence of events from inquiry to ironic triumph")

    steps = [
        (
            "01",
            "The Rental Inquiry",
            "The speaker finds a promising rental advertisement in London. Pleased with the reasonable rent and location, he calls the landlady from a public telephone booth to secure accommodation."
        ),
        (
            "02",
            "Pre-emptive Disclosure",
            "Knowing the prevalence of discrimination, the speaker chooses to avoid wasted travel. He politely and frankly discloses his racial identity, stating plainly: 'I am African.'"
        ),
        (
            "03",
            "The Abrupt Question",
            "A tense, uncomfortable silence follows. Instead of asking about his profession or character, the landlady coldly demands clarification on his skin shade: 'HOW DARK?'"
        ),
        (
            "04",
            "Absurd Classification",
            "Failing to understand the nuance, the landlady bluntly asks if he is 'light or very dark.' The speaker humorously categorizes his complexion as 'West African sepia.'"
        ),
        (
            "05",
            "The Ironic Reversal",
            "The speaker playfully explains that his face is sepia, but his palms and feet are lighter. He begs her to inspect him before hanging up—turning racial prejudice into total absurdity."
        )
    ]

    col_w = Inches(2.2)
    gap = Inches(0.18)
    start_x = Inches(0.8)
    step_y = Inches(1.9)
    step_h = Inches(3.9)

    for i, (num_str, title_txt, desc_txt) in enumerate(steps):
        cur_x = start_x + i * (col_w + gap)

        add_card(slide3, cur_x, step_y, col_w, step_h, fill_color=COLOR_CARD, border_color=COLOR_CARD_BORDER, border_width=Pt(0.75))

        # Top accent bar for stage
        sbar = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, cur_x, step_y, col_w, Inches(0.06))
        sbar.fill.solid()
        sbar.fill.fore_color.rgb = COLOR_ACCENT
        sbar.line.color.rgb = COLOR_ACCENT
        sbar.line.width = Pt(0)

        tb = slide3.shapes.add_textbox(cur_x + Inches(0.15), step_y + Inches(0.18), col_w - Inches(0.3), step_h - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_num = tf.paragraphs[0]
        p_num.text = num_str
        p_num.font.name = FONT_TITLE
        p_num.font.size = Pt(20)
        p_num.font.bold = True
        p_num.font.color.rgb = COLOR_ACCENT
        p_num.space_after = Pt(4)

        p_st = tf.add_paragraph()
        p_st.text = title_txt
        p_st.font.name = FONT_TITLE
        p_st.font.size = Pt(13)
        p_st.font.bold = True
        p_st.font.color.rgb = COLOR_TEXT_PRIMARY
        p_st.space_after = Pt(8)

        p_sd = tf.add_paragraph()
        p_sd.text = desc_txt
        p_sd.font.name = FONT_BODY
        p_sd.font.size = Pt(11)
        p_sd.font.color.rgb = COLOR_TEXT_BODY

    # Bottom takeaway bar
    bot_y = Inches(6.0)
    add_card(slide3, Inches(0.8), bot_y, Inches(11.733), Inches(0.68), fill_color=COLOR_ACCENT_TINT, border_color=COLOR_CARD_BORDER, border_width=Pt(0.75))

    tb_bot = slide3.shapes.add_textbox(Inches(1.0), bot_y + Inches(0.12), Inches(11.333), Inches(0.48))
    tf_bot = tb_bot.text_frame
    tf_bot.word_wrap = True
    tf_bot.margin_left = tf_bot.margin_top = tf_bot.margin_right = tf_bot.margin_bottom = 0
    p_bot = tf_bot.paragraphs[0]
    p_bot.text = "Key Narrative Movement: The speaker shifts from polite accommodation to intellectual mastery, using humor to expose the landlady's narrow bigotry."
    p_bot.font.name = FONT_BODY
    p_bot.font.size = Pt(12)
    p_bot.font.italic = True
    p_bot.font.color.rgb = COLOR_ACCENT

    add_footer(slide3, 3)
    set_notes(slide3,
        "Let's trace the narrative sequence of the poem step by step. The entire poem takes place within a single phone conversation.\n\n"
        "In Stage 1, our speaker is searching for a place to rent in London. He finds an advert with a reasonable price and an indifferent location, "
        "and calls the landlady from a public red telephone booth.\n\n"
        "In Stage 2, because racial discrimination was so rampant in 1960s Britain, the speaker decides to pre-empt a wasted journey. "
        "He politely confesses: 'Madam, I am African.' He chooses complete honesty.\n\n"
        "In Stage 3, there is a heavy, awkward silence on the other end of the line. The landlady reveals her prejudice when she asks abruptly: "
        "'HOW DARK?' She completely ignores his manners or suitability as a tenant.\n\n"
        "In Stage 4, she repeats the question, asking whether he is light or very dark. Instead of responding with anger, the speaker responds with wit. "
        "He compares his complexion to artistic shades like 'West African sepia,' mocking her attempt to categorize human beings like paint samples.\n\n"
        "Finally, in Stage 5, the speaker explains that his face is brunette, but the palms of his hands and soles of his feet are lighter, "
        "and pleads with her to inspect his backside before hanging up. Through irony, he strips the landlady of her power and exposes the absurdity of racism. "
        "Now, let's explore the major themes."
    )

    # =========================================================================
    # SLIDE 4: MAJOR THEMES
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4)
    add_header(slide4, "Thematic Analysis", "Major Themes & Social Message", "Exploring racial prejudice, human dignity, and everyday hypocrisy")

    themes = [
        (
            "1. Racial Discrimination",
            "The speaker encounters prejudice during an ordinary housing inquiry. Soyinka illustrates how racism is not just an abstract ideology, but a concrete daily obstacle that denies people equal access to basic necessities like shelter."
        ),
        (
            "2. Human Dignity & Self-Respect",
            "Despite being subjected to intrusive and insulting questions about his physical appearance, the speaker never loses his composure. He maintains his self-worth and moral dignity, refusing to be reduced to a stereotype."
        ),
        (
            "3. Prejudice & Polite Hypocrisy",
            "The landlady adopts a veneer of upper-class politeness, but her polite manner conceals deep-seated bias. Soyinka critiques the hypocrisy of 'respectable' society that practices polite discrimination."
        ),
        (
            "4. Identity & Skin Colour",
            "The poem attacks the irrational obsession with classifying human worth by skin tone. Soyinka demonstrates that judging a person's intelligence, character, or value by pigmentation is fundamentally absurd."
        )
    ]

    card_w = Inches(5.72)
    card_h = Inches(1.9)
    gap_x = Inches(0.29)
    gap_y = Inches(0.2)

    positions = [
        (Inches(0.8), Inches(1.9)),
        (Inches(0.8) + card_w + gap_x, Inches(1.9)),
        (Inches(0.8), Inches(1.9) + card_h + gap_y),
        (Inches(0.8) + card_w + gap_x, Inches(1.9) + card_h + gap_y),
    ]

    for (title_txt, desc_txt), (x_pos, y_pos) in zip(themes, positions):
        add_card(slide4, x_pos, y_pos, card_w, card_h, fill_color=COLOR_CARD, border_color=COLOR_CARD_BORDER, border_width=Pt(0.75))

        # Left subtle accent bar
        tbar = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, x_pos, y_pos, Inches(0.08), card_h)
        tbar.fill.solid()
        tbar.fill.fore_color.rgb = COLOR_ACCENT
        tbar.line.color.rgb = COLOR_ACCENT
        tbar.line.width = Pt(0)

        tb = slide4.shapes.add_textbox(x_pos + Inches(0.24), y_pos + Inches(0.16), card_w - Inches(0.4), card_h - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_tt = tf.paragraphs[0]
        p_tt.text = title_txt
        p_tt.font.name = FONT_TITLE
        p_tt.font.size = Pt(14)
        p_tt.font.bold = True
        p_tt.font.color.rgb = COLOR_ACCENT
        p_tt.space_after = Pt(4)

        p_td = tf.add_paragraph()
        p_td.text = desc_txt
        p_td.font.name = FONT_BODY
        p_td.font.size = Pt(11.5)
        p_td.font.color.rgb = COLOR_TEXT_BODY

    # Concluding thematic message banner
    tmsg_y = Inches(6.05)
    add_card(slide4, Inches(0.8), tmsg_y, Inches(11.733), Inches(0.65), fill_color=COLOR_ACCENT_TINT, border_color=COLOR_CARD_BORDER, border_width=Pt(0.75))

    tb_tm = slide4.shapes.add_textbox(Inches(1.0), tmsg_y + Inches(0.14), Inches(11.333), Inches(0.45))
    tf_tm = tb_tm.text_frame
    tf_tm.word_wrap = True
    tf_tm.margin_left = tf_tm.margin_top = tf_tm.margin_right = tf_tm.margin_bottom = 0
    p_tm = tf_tm.paragraphs[0]
    p_tm.text = "Concluding Insight: A person's character, intellect, and humanity matter far more than their skin colour."
    p_tm.font.name = FONT_TITLE
    p_tm.font.size = Pt(12.5)
    p_tm.font.bold = True
    p_tm.font.color.rgb = COLOR_ACCENT

    add_footer(slide4, 4)
    set_notes(slide4,
        "Moving on to the major themes of the poem, Soyinka weaves several profound social realities into this short dialogue.\n\n"
        "First, Racial Discrimination: Soyinka focuses on institutional and everyday racism. It is not overt physical violence, "
        "but the quiet rejection that prevents a qualified, polite student from having a roof over his head.\n\n"
        "Second is Human Dignity: the speaker refuses to let the landlady's degrading inquiry strip him of his self-respect. "
        "Instead of pleading or showing distress, he maintains intellectual superiority and self-command.\n\n"
        "Third is Prejudice and Polite Hypocrisy: the landlady hides behind polite diction, but beneath her voice lies cold racial bias. "
        "Soyinka reveals how British society often maintained polite appearances while practicing systemic exclusion.\n\n"
        "Fourth is Identity and Skin Colour: the poem ridicules the obsession with shade and melanin. "
        "Skin colour is an accident of biology, not an indicator of integrity, culture, or worth.\n\n"
        "Ultimately, Soyinka reminds us that character and shared humanity matter far more than outward physical markers. "
        "Next, let’s see how Soyinka communicates this through literary craft."
    )

    # =========================================================================
    # SLIDE 5: LITERARY DEVICES
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5)
    add_header(slide5, "Poetic Craft", "Literary Devices & Techniques", "How Soyinka uses irony, satire, and imagery to expose prejudice")

    left_x = Inches(0.8)
    col_w = Inches(5.72)
    dev_h = Inches(1.22)
    dgap = Inches(0.13)

    left_devices = [
        (
            "Irony",
            "The speaker uses exaggerated courtesy and intellectual vocabulary to address an ignorant question. This exposes the deep irony that the African caller is far more cultured and civilized than the British landlady."
        ),
        (
            "Satire",
            "Soyinka relies on biting humor and mockery to attack racism. Rather than delivering an angry lecture, he ridicules the landlady's obsession with skin tones, stripping bigotry of its authority."
        ),
        (
            "Imagery",
            "Vivid sensory details bring the scene alive: the 'red telephone booth,' 'red omnibus,' and shades like 'West African sepia' and 'peroxide blonde.' These images contrast external colours with internal moral reality."
        )
    ]

    dy = Inches(1.9)
    for d_name, d_desc in left_devices:
        add_card(slide5, left_x, dy, col_w, dev_h, fill_color=COLOR_CARD, border_color=COLOR_CARD_BORDER, border_width=Pt(0.75))

        dbar = slide5.shapes.add_shape(MSO_SHAPE.RECTANGLE, left_x, dy, Inches(0.08), dev_h)
        dbar.fill.solid()
        dbar.fill.fore_color.rgb = COLOR_ACCENT
        dbar.line.color.rgb = COLOR_ACCENT
        dbar.line.width = Pt(0)

        tb = slide5.shapes.add_textbox(left_x + Inches(0.22), dy + Inches(0.1), col_w - Inches(0.35), dev_h - Inches(0.18))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p1 = tf.paragraphs[0]
        p1.text = d_name
        p1.font.name = FONT_TITLE
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_ACCENT
        p1.space_after = Pt(2)

        p2 = tf.add_paragraph()
        p2.text = d_desc
        p2.font.name = FONT_BODY
        p2.font.size = Pt(11)
        p2.font.color.rgb = COLOR_TEXT_BODY

        dy += dev_h + dgap

    # Right column: 2 devices + Critical Synthesis Box
    right_x = Inches(0.8) + col_w + Inches(0.29)
    right_devices = [
        (
            "Contrast",
            "A sharp contrast is drawn between the speaker's calm, articulate eloquence and the landlady's blunt, insensitive questions. This subverts colonial stereotypes about who possesses refinement."
        ),
        (
            "Dialogue",
            "The poem is structured as an authentic, overheard telephone dialogue. This gives the interaction immediate dramatic tension and realism, making the audience feel the sting of discrimination firsthand."
        )
    ]

    ry = Inches(1.9)
    for d_name, d_desc in right_devices:
        add_card(slide5, right_x, ry, col_w, dev_h, fill_color=COLOR_CARD, border_color=COLOR_CARD_BORDER, border_width=Pt(0.75))

        dbar = slide5.shapes.add_shape(MSO_SHAPE.RECTANGLE, right_x, ry, Inches(0.08), dev_h)
        dbar.fill.solid()
        dbar.fill.fore_color.rgb = COLOR_ACCENT
        dbar.line.color.rgb = COLOR_ACCENT
        dbar.line.width = Pt(0)

        tb = slide5.shapes.add_textbox(right_x + Inches(0.22), ry + Inches(0.1), col_w - Inches(0.35), dev_h - Inches(0.18))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p1 = tf.paragraphs[0]
        p1.text = d_name
        p1.font.name = FONT_TITLE
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_ACCENT
        p1.space_after = Pt(2)

        p2 = tf.add_paragraph()
        p2.text = d_desc
        p2.font.name = FONT_BODY
        p2.font.size = Pt(11)
        p2.font.color.rgb = COLOR_TEXT_BODY

        ry += dev_h + dgap

    # Critical Synthesis Box
    add_card(slide5, right_x, ry, col_w, Inches(1.22), fill_color=COLOR_ACCENT_TINT, border_color=COLOR_CARD_BORDER, border_width=Pt(0.75))

    tb_sy = slide5.shapes.add_textbox(right_x + Inches(0.22), ry + Inches(0.12), col_w - Inches(0.4), Inches(1.0))
    tf_sy = tb_sy.text_frame
    tf_sy.word_wrap = True
    tf_sy.margin_left = tf_sy.margin_top = tf_sy.margin_right = tf_sy.margin_bottom = 0

    ps1 = tf_sy.paragraphs[0]
    ps1.text = "How Poetic Craft Serves the Social Message"
    ps1.font.name = FONT_TITLE
    ps1.font.size = Pt(12)
    ps1.font.bold = True
    ps1.font.color.rgb = COLOR_ACCENT
    ps1.space_after = Pt(3)

    ps2 = tf_sy.add_paragraph()
    ps2.text = "By fusing realistic dialogue with sharp irony and satire, Soyinka avoids preaching. Instead, he empowers the speaker to intellectually dismantle prejudice with wit and composure."
    ps2.font.name = FONT_BODY
    ps2.font.size = Pt(10.5)
    ps2.font.italic = True
    ps2.font.color.rgb = COLOR_TEXT_BODY

    # Bottom summary rule note
    bs_y = Inches(6.05)
    add_card(slide5, Inches(0.8), bs_y, Inches(11.733), Inches(0.65), fill_color=COLOR_CARD, border_color=COLOR_CARD_BORDER, border_width=Pt(0.75))

    tb_bs = slide5.shapes.add_textbox(Inches(1.0), bs_y + Inches(0.12), Inches(11.333), Inches(0.45))
    tf_bs = tb_bs.text_frame
    tf_bs.word_wrap = True
    tf_bs.margin_left = tf_bs.margin_top = tf_bs.margin_right = tf_bs.margin_bottom = 0
    pbs = tf_bs.paragraphs[0]
    pbs.text = "Literary Impact: Soyinka turns a humiliating situation into an intellectual victory, proving that humor is a powerful weapon against ignorance."
    pbs.font.name = FONT_BODY
    pbs.font.size = Pt(11.5)
    pbs.font.bold = True
    pbs.font.color.rgb = COLOR_ACCENT

    add_footer(slide5, 5)
    set_notes(slide5,
        "Now let’s explore how Soyinka uses poetic craft to communicate his message. Rather than writing a simple protest poem, "
        "he uses sophisticated literary techniques.\n\n"
        "First is Irony: the speaker speaks with great refinement and politeness, while the British landlady—who considers herself civilized—"
        "behaves with crude prejudice. The irony completely flips the supposed social hierarchy.\n\n"
        "Second is Satire: Soyinka uses sharp humor and ridicule. When the speaker asks the landlady to inspect the palms of his hands "
        "and the soles of his feet, he uses satire to make her obsession with skin color look completely ridiculous.\n\n"
        "Third is Imagery: the poem is filled with sensory details—the red telephone booth, the red double-decker omnibus, the smell of rancid breath, "
        "and artistic shades like 'sepia' and 'peroxide blonde.' This makes the encounter feel physical and grounded.\n\n"
        "Fourth is Contrast: we see a deliberate contrast between the speaker's articulate calm and the landlady's blunt interrogation. "
        "This subverts colonial assumptions about intellect and refinement.\n\n"
        "And fifth, Dialogue: framing the poem as a phone conversation gives it dramatic immediacy and realism. "
        "Together, these devices allow Soyinka to expose prejudice without preaching, turning wit into an instrument of resistance. "
        "Let us now conclude our presentation."
    )

    # =========================================================================
    # SLIDE 6: CONCLUSION
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6)
    add_header(slide6, "Summary & Reflection", "Conclusion: The Enduring Message", "Why Soyinka's poem continues to resonate in contemporary society")

    conc_items = [
        (
            "Racism Reduces Humanity",
            "Prejudice strips individuals of their intellect, character, and uniqueness, reducing complex human souls to surface physical traits."
        ),
        (
            "Discrimination Denies Dignity",
            "Everyday bias creates unjust barriers, denying people basic human rights such as equal respect, housing, and social inclusion."
        ),
        (
            "Wit as Moral Resistance",
            "Humour and irony can effectively disarm bigotry, exposing the irrationality and foolishness of racial intolerance."
        ),
        (
            "Universal Relevance",
            "Written in the 1960s, the poem remains an urgent reminder for modern societies to judge individuals by character, not appearance."
        )
    ]

    c_w = Inches(2.78)
    c_gap = Inches(0.2)
    c_y = Inches(1.9)
    c_h = Inches(2.2)

    for i, (ctitle, cdesc) in enumerate(conc_items):
        cx = Inches(0.8) + i * (c_w + c_gap)
        add_card(slide6, cx, c_y, c_w, c_h, fill_color=COLOR_CARD, border_color=COLOR_CARD_BORDER, border_width=Pt(0.75))

        # Top accent bar
        cbar = slide6.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx, c_y, c_w, Inches(0.06))
        cbar.fill.solid()
        cbar.fill.fore_color.rgb = COLOR_ACCENT
        cbar.line.color.rgb = COLOR_ACCENT
        cbar.line.width = Pt(0)

        tb = slide6.shapes.add_textbox(cx + Inches(0.18), c_y + Inches(0.16), c_w - Inches(0.36), c_h - Inches(0.28))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p1 = tf.paragraphs[0]
        p1.text = ctitle
        p1.font.name = FONT_TITLE
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_ACCENT
        p1.space_after = Pt(6)

        p2 = tf.add_paragraph()
        p2.text = cdesc
        p2.font.name = FONT_BODY
        p2.font.size = Pt(11)
        p2.font.color.rgb = COLOR_TEXT_BODY

    # Prominent Concluding Statement Card
    q_y = Inches(4.35)
    add_card(slide6, Inches(0.8), q_y, Inches(11.733), Inches(1.6), fill_color=COLOR_CARD, border_color=COLOR_ACCENT, border_width=Pt(1.5))

    tb_q = slide6.shapes.add_textbox(Inches(1.2), q_y + Inches(0.18), Inches(10.933), Inches(1.25))
    tf_q = tb_q.text_frame
    tf_q.word_wrap = True
    tf_q.margin_left = tf_q.margin_top = tf_q.margin_right = tf_q.margin_bottom = 0

    pq_eyebrow = tf_q.paragraphs[0]
    pq_eyebrow.text = "CORE SEMINAR TAKEAWAY"
    pq_eyebrow.alignment = PP_ALIGN.CENTER
    pq_eyebrow.font.name = FONT_BODY
    pq_eyebrow.font.size = Pt(10)
    pq_eyebrow.font.bold = True
    pq_eyebrow.font.color.rgb = COLOR_ACCENT
    pq_eyebrow.space_after = Pt(6)

    pq = tf_q.add_paragraph()
    pq.text = "“Every person deserves equal dignity and respect, regardless of skin colour.”"
    pq.alignment = PP_ALIGN.CENTER
    pq.font.name = FONT_TITLE
    pq.font.size = Pt(18)
    pq.font.bold = True
    pq.font.italic = True
    pq.font.color.rgb = COLOR_TEXT_PRIMARY
    pq.space_after = Pt(5)

    pq_sub = tf_q.add_paragraph()
    pq_sub.text = "Human dignity has no colour, and equality is a right, not a privilege."
    pq_sub.alignment = PP_ALIGN.CENTER
    pq_sub.font.name = FONT_BODY
    pq_sub.font.size = Pt(12)
    pq_sub.font.color.rgb = COLOR_TEXT_MUTED

    # Bottom sign-off
    tb_end = slide6.shapes.add_textbox(Inches(0.8), Inches(6.12), Inches(11.733), Inches(0.45))
    tf_end = tb_end.text_frame
    tf_end.word_wrap = True
    tf_end.margin_left = tf_end.margin_top = tf_end.margin_right = tf_end.margin_bottom = 0
    pe = tf_end.paragraphs[0]
    pe.text = "Thank you for your attention • Questions and seminar discussion are warmly welcomed."
    pe.alignment = PP_ALIGN.CENTER
    pe.font.name = FONT_BODY
    pe.font.size = Pt(11.5)
    pe.font.italic = True
    pe.font.color.rgb = COLOR_TEXT_MUTED

    add_footer(slide6, 6)
    set_notes(slide6,
        "To conclude our seminar, Wole Soyinka’s 'Telephone Conversation' remains an enduring classic because it speaks directly to our shared humanity.\n\n"
        "It exposes the tragedy of prejudice: that someone can be judged entirely by their skin color before anything is known about their character, "
        "intellect, or decency. Yet the poem is fundamentally empowering. The speaker refuses to accept the role of a humiliated victim; "
        "instead, he uses humour and superior intellect to reveal the absurdity of his oppressor.\n\n"
        "As we reflect on this poem in our English seminar today, its moral truth is clear and undeniable: "
        "'Every person deserves equal dignity and respect, regardless of skin colour.' Human dignity has no colour, and equality is a right, "
        "not a privilege.\n\n"
        "Thank you very much for your time and attention. I would now be delighted to open the floor to any questions, thoughts, or discussion from the seminar."
    )

    output_path = "Telephone_Conversation_Seminar.pptx"
    prs.save(output_path)
    print(f"Presentation successfully created and saved at: {output_path}")

if __name__ == "__main__":
    create_deck()
